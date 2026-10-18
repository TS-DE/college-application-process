"""文件处理工具：UUID 重命名保存 + 文本提取（PDF / DOCX / TXT）。"""
import os
import uuid

from app.config import settings

ALLOWED_EXTS = {".pdf", ".txt", ".docx"}
MAX_FILE_SIZE = 20 * 1024 * 1024  # 20MB

# 0.12.1：用户画像报告的**独立**白名单与目录。
# ★ 与知识库（ALLOWED_EXTS / KNOWLEDGE_DIR）彻底分开——用户的私人报告
#   绝不能和管理员的公共政策文件混在一个目录、走同一套上传链路。
REPORT_ALLOWED_EXTS = {".pdf", ".txt", ".docx", ".md", ".png", ".jpg", ".jpeg"}
REPORT_MAX_SIZE = 10 * 1024 * 1024  # 报告通常几页，10MB 足够

TYPE_MAP = {
    ".pdf": "pdf",
    ".txt": "txt",
    ".docx": "docx",
}


def ensure_dir() -> str:
    os.makedirs(settings.KNOWLEDGE_DIR, exist_ok=True)
    return settings.KNOWLEDGE_DIR


def ensure_report_dir() -> str:
    """用户报告的私有落盘目录（0.12.1，与知识库目录物理隔离）。"""
    path = os.path.join(settings.UPLOAD_ROOT, "reports")
    os.makedirs(path, exist_ok=True)
    return path


def save_report_file(content: bytes, original_filename: str) -> tuple[str, str]:
    """保存用户报告到独立目录，返回 (stored_filename, file_path)。

    文件会在文本提取完成后被**立即删除**（见 services/profile/report_reader.py），
    磁盘上不留私人报告副本。
    """
    ext = ext_of(original_filename)
    safe_original = os.path.basename(original_filename).replace(" ", "_")
    stored = f"{uuid.uuid4().hex}_{safe_original}"
    if ext and not stored.endswith(ext):
        stored += ext
    path = os.path.join(ensure_report_dir(), stored)
    with open(path, "wb") as f:
        f.write(content)
    return stored, path


def is_report_allowed(filename: str) -> bool:
    return ext_of(filename) in REPORT_ALLOWED_EXTS


def ext_of(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def is_allowed(filename: str) -> bool:
    return ext_of(filename) in ALLOWED_EXTS


def save_upload_file(content: bytes, original_filename: str) -> tuple[str, str]:
    """保存上传内容，返回 (stored_filename, file_path)。

    stored_filename 使用 UUID 重命名，避免冲突与路径遍历。
    """
    ext = ext_of(original_filename)
    safe_original = os.path.basename(original_filename).replace(" ", "_")
    stored = f"{uuid.uuid4().hex}_{safe_original}"
    if ext and not stored.endswith(ext):
        stored += ext
    directory = ensure_dir()
    path = os.path.join(directory, stored)
    with open(path, "wb") as f:
        f.write(content)
    return stored, path


def file_type_of(filename: str) -> str:
    return TYPE_MAP.get(ext_of(filename), ext_of(filename).lstrip(".") or "unknown")


def extract_text(file_path: str, file_type: str) -> tuple[str, bool]:
    """按类型提取纯文本。

    0.10.0：返回值从 `str` 改为 `(文本, 是否用了 OCR)`。
    扫描件 PDF 直接提取只能拿到页码（每页 1 字符），此时自动走 OCR 兜底；
    调用方需要把 `ocr_used` 回写到 knowledge_files 表，便于排查与统计。

    :return: (text, ocr_used)
    """
    ext = (file_type or ext_of(file_path)).lower().lstrip(".")
    if ext == "pdf":
        text = _extract_pdf(file_path)
        return _maybe_ocr(file_path, text)
    if ext == "docx":
        return _extract_docx(file_path), False
    return _extract_txt(file_path), False


def _maybe_ocr(file_path: str, text: str) -> tuple[str, bool]:
    """PDF 提取结果太薄时尝试 OCR 兜底（0.10.0）。

    只有 PDF 会走到这里；OCR 未开启 / 非扫描件 / OCR 失败时原样返回原文。
    """
    if not text or not text.strip():
        # 完全提不到文字：仍要走一遍 OCR 判定（扫描件判定本身就看字符数）
        pass
    try:
        from app.utils.ocr_utils import OCRExtractor

        ocr_text, used = OCRExtractor().extract_text_with_ocr_fallback(file_path)
        if used and ocr_text:
            return ocr_text, True
    except Exception:  # noqa: BLE001 OCR 模块自身异常不能影响上传主流程
        pass
    return text, False


def _extract_txt(file_path: str) -> str:
    for enc in ("utf-8", "gbk", "utf-16"):
        try:
            with open(file_path, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _extract_pdf(file_path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(file_path)
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(file_path: str) -> str:
    from docx import Document

    doc = Document(file_path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def split_text(text: str, size: int | None = None, overlap: int | None = None) -> list[str]:
    """按字数切分，默认 500 字一片、重叠 50 字。"""
    size = size or settings.CHUNK_SIZE
    overlap = overlap or settings.CHUNK_OVERLAP
    text = (text or "").strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]
    step = max(1, size - overlap)
    chunks = []
    start = 0
    while start < len(text):
        chunks.append(text[start : start + size])
        start += step
    return chunks
